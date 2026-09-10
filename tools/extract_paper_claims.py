"""Extract the manuscript's table values into a machine-readable snapshot.

The repository does not publish the manuscript, so verification needs the
numbers it asserts in a form a reviewer can diff against the evidence. This
writes only table cells plus the manuscript's SHA-256 - no prose, no author
information, no affiliation.

Run from the paper working directory:

    python tools/extract_paper_claims.py --docx <manuscript.docx> \\
        --out paper_claims.json
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph

CAPTION = re.compile(r"^(续?表)\s*(\d+)\s")


def _blocks(document: Document):
    """Yield paragraphs and tables in document order."""
    body = document.element.body
    paragraphs = {p._p: p for p in document.paragraphs}
    tables = {t._element: t for t in document.tables}
    for node in body:
        if node in paragraphs:
            yield paragraphs[node]
        elif node in tables:
            yield tables[node]


def logical_tables(document: Document) -> dict[int, list[list[str]]]:
    """Group tables by caption number, merging any 续表 continuation.

    A table split across a page becomes two table objects with a repeated
    header. Keying on the caption rather than the object index keeps the
    logical table intact and prevents a continuation from being missed.
    """
    collected: dict[int, list[list[str]]] = {}
    pending: tuple[int, bool] | None = None
    for block in _blocks(document):
        if isinstance(block, Paragraph):
            match = CAPTION.match(block.text.strip())
            if match:
                pending = (int(match.group(2)), match.group(1) == "续表")
            continue
        if not isinstance(block, Table) or pending is None:
            continue
        number, is_continuation = pending
        rows = [[cell.text.strip() for cell in row.cells] for row in block.rows]
        if is_continuation and number in collected:
            # Drop the repeated header row of the continuation.
            collected[number].extend(rows[1:])
        else:
            collected.setdefault(number, []).extend(rows)
        pending = None
    return collected


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--docx", required=True, type=Path)
    parser.add_argument("--out", required=True, type=Path)
    parser.add_argument("--title", default="residual-weight explainable IDS manuscript")
    args = parser.parse_args()

    document = Document(str(args.docx))
    tables = logical_tables(document)
    payload = {
        "source": {
            "title": args.title,
            "sha256": hashlib.sha256(args.docx.read_bytes()).hexdigest(),
            "note": "Table cells only. No manuscript prose or author data.",
        },
        "tables": {
            str(number): tables[number] for number in sorted(tables) if number <= 8
        },
    }
    args.out.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    counts = {n: len(rows) for n, rows in payload["tables"].items()}
    print(json.dumps({"tables": counts, "sha256": payload["source"]["sha256"][:16]}))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
